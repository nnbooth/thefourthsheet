#!/usr/bin/env python3
"""
make_calls_templates.py — made the first version of the Calls report's templates (22 Oct 2026). One-off.

    python3 tools/archive/make_calls_templates.py

Writes assets/exports/templates/calls.xlsx and calls.pptx: empty layouts with {{slots}} where the data goes, no
numbers in them. The site fills them in the visitor's browser when they click Export (assets/exports/fill.js, from
data/reports/calls.json). From now on the template files are the source: open them in Excel or PowerPoint and edit
them there. Running this again overwrites those edits.

The slots (assets/exports/fill.js explains them in full):
  {{answer}}            text
  {{n:w.calls}}         a number (Excel cells)
  {{months.label}}      a row repeated for every month (Excel; the last row on its sheet)
  chart alt text "chart:hours"   a PowerPoint chart whose data is replaced with the "hours" chart
"""
import sys, pathlib  # noqa: E401  (next 2 lines: make every stage folder importable; see tools/README.md)
_T = next(p for p in pathlib.Path(__file__).resolve().parents if p.name == "tools")
sys.path[:0] = [d for d in (str(x) for x in sorted(_T.iterdir()) if x.is_dir() and (x.name[0].isdigit() or x.name in ("shared", "setup"))) if d not in sys.path]

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import AreaChart, BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, PatternFill
from openpyxl.worksheet.properties import PageSetupProperties
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

import exportkit as ek
import pptkit

Font = ek.xl_font
REPO = Path(__file__).resolve().parents[2]
OUT = REPO / "assets" / "exports" / "templates"
GREEN, INK, MUTED, BRAND = "2F7A5D", "25342A", "5F6F63", "2F7A5D"
INT, PCT, SEC = '#,##0;[Red](#,##0);"-"', '0.0%;[Red](0.0%);"-"', '0" sec"'
SUB = "{{business}} · SAMPLE DATA (invented) · The Fourth Sheet"
STATUS = "Period: {{period_label}} · Data status: {{status_short}} · {{generated}}"


# ------------------------------------------------------------------ Excel

def title(ws, text):
    ws["A1"], ws["A2"], ws["A3"] = text, SUB, STATUS
    ws["A1"].font = Font(bold=True, size=14, color=INK)
    ws["A2"].font = ws["A3"].font = Font(italic=True, size=9, color=MUTED)


def head(ws, row, cells, left=(1,)):
    for c, h in enumerate(cells, 1):
        cell = ws.cell(row, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.alignment = Alignment(horizontal="left" if c in left else "right", vertical="center", wrap_text=True, indent=1 if c in left and c > 1 else 0)


def para(ws, row, text, bold=False, cols=4, height=32, colour=INK):
    ws.cell(row, 1, text).font = Font(bold=bold, color=colour)
    ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    ws.row_dimensions[row].height = height


def printable(ws, one_page=False):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, (1 if one_page else 0)
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.5
    ws.page_margins.top, ws.page_margins.bottom = 0.6, 0.6
    ws.print_title_rows = None
    ws.oddFooter.left.text = "{{business}} · {{question}}"
    ws.oddFooter.right.text = "Page &P of &N"
    ws.oddHeader.right.text = "{{generated}}"
    for part in (ws.oddFooter.left, ws.oddFooter.right, ws.oddHeader.right):
        part.size, part.font = 8, f"{ek.XL_FONT},Regular"


def xlsx():
    wb = ek.xl_default_font(Workbook())
    rep = wb.active
    rep.title = "Report"
    wk, hrs, data = wb.create_sheet("Workings"), wb.create_sheet("Hours"), wb.create_sheet("Data")

    # Report: the answer, the headline numbers (each a formula on Workings), what it shows, the charts
    title(rep, "{{question}}")
    para(rep, 5, "In short: {{answer}}", bold=True)
    para(rep, 6, "{{intro}}")
    rep["A7"] = "{{part_note}}"
    rep["A7"].font = Font(italic=True, size=9, color=MUTED)
    head(rep, 8, ["Headline number", "Value", "Compared with"], left=(1, 3))
    for i, (ref, fmt) in enumerate((("B8", INT), ("B11", PCT), ("B12", PCT), ("B18", SEC)), 1):
        rep.cell(8 + i, 1, f"{{{{kpi{i}.label}}}}")
        rep.cell(8 + i, 2, f"=Workings!{ref}").number_format = fmt
        rep.cell(8 + i, 3, f"{{{{kpi{i}.sub}}}}").font = Font(color=MUTED)
        rep.cell(8 + i, 3).alignment = Alignment(indent=1)
    rep["A14"] = "What it shows"
    rep["A14"].font = Font(bold=True, color=INK)
    para(rep, 15, "{{shows}}", height=48)
    rep["A17"] = "What you'd do about it"
    rep["A17"].font = Font(bold=True, color=INK)
    para(rep, 18, "{{action}}", height=48)
    rep["A20"] = "Every time of day is on the Hours sheet and every month on the Data sheet; how each headline number is worked out is on Workings."
    rep["A20"].font = Font(italic=True, size=9, color=MUTED)
    for col, w in zip("ABCD", (44, 14, 26, 4)):
        rep.column_dimensions[col].width = w

    # Workings: the counts as values, every rate a formula
    title(wk, "How each number is worked out")
    wk["A5"] = "Calls for help, {{cur}} and {{prev}}"
    wk["A6"] = "Answered ÷ calls while the line was staffed"
    head(wk, 7, ["", "{{cur}}", "{{prev}}"])
    for r, (lab, key) in enumerate((("Calls", "calls"), ("Calls while staffed", "staffed"), ("Answered", "answered")), 8):
        wk.cell(r, 1, lab)
        wk.cell(r, 2, f"{{{{n:w.{key}}}}}").number_format = INT
        wk.cell(r, 3, f"{{{{n:w.{key}_prev}}}}").number_format = INT
    wk["A11"], wk["A12"] = "Answered while staffed", "Calls outside staffed hours"
    for c in "BC":
        wk[f"{c}11"] = f'=IF({c}9=0,"",{c}10/{c}9)'
        wk[f"{c}12"] = f'=IF({c}8=0,"",({c}8-{c}9)/{c}8)'
        wk[f"{c}11"].number_format = wk[f"{c}12"].number_format = PCT
    wk["A14"] = "Median wait when answered, {{cur}} and {{prev}}"
    wk["A15"] = "The middle wait when every answered call is lined up from shortest to longest"
    head(wk, 16, ["", "{{cur}}", "{{prev}}"])
    wk["A17"], wk["B17"], wk["C17"] = "Answered calls (while staffed)", "{{n:w.median_n}}", "{{n:w.median_n_prev}}"
    wk["A18"], wk["B18"], wk["C18"] = "Median wait (seconds, rounded half up)", "{{n:w.median}}", "{{n:w.median_prev}}"
    para(wk, 19, "{{w.median_note}}", cols=3)
    for a in ("A5", "A14"):
        wk[a].font = Font(bold=True, color=INK)
    for a in ("A6", "A15"):
        wk[a].font = Font(italic=True, size=9, color=MUTED)
    wk.column_dimensions["A"].width = 44
    wk.column_dimensions["B"].width = wk.column_dimensions["C"].width = 14

    # Hours: one row per time of day (the row repeats for each one)
    title(hrs, "When the calls come in, {{period_label}}")
    head(hrs, 5, ["Time of day", "Calls", "Answered", "Answered %"])
    hrs["A6"], hrs["B6"], hrs["C6"], hrs["D6"] = "{{hours.label}}", "{{n:hours.calls}}", "{{n:hours.answered}}", '=IF(B6=0,0,C6/B6)'
    hrs["B6"].number_format = hrs["C6"].number_format = INT
    hrs["D6"].number_format = PCT
    for col, w in zip("ABCD", (22, 12, 12, 14)):
        hrs.column_dimensions[col].width = w

    # Data: one row per month (the row repeats for each one)
    title(data, "Calls for help by month")
    head(data, 5, ["Month", "Status", "Calls", "While staffed", "Answered", "Answered while staffed"])
    for c, v in zip("ABCDE", ("{{months.label}}", "{{months.status}}", "{{n:months.calls}}", "{{n:months.staffed}}", "{{n:months.answered}}")):
        data[f"{c}6"] = v
        if c > "B":
            data[f"{c}6"].number_format = INT
    data["F6"] = '=IF(D6=0,"",E6/D6)'
    data["F6"].number_format = PCT
    for col, w in zip("ABCDEF", (10, 12, 10, 14, 12, 22)):
        data.column_dimensions[col].width = w

    # charts on Report, drawn from the repeating rows (the fill stretches their ranges to every row)
    bars = BarChart()
    bars.type, bars.title, bars.legend, bars.gapWidth = "bar", "Calls by time of day", None, 60
    bars.add_data(Reference(hrs, min_col=2, min_row=6, max_row=6), titles_from_data=False)
    bars.set_categories(Reference(hrs, min_col=1, min_row=6, max_row=6))
    bars.x_axis.scaling.orientation = "maxMin"          # first time of day at the top, like the table
    bars.y_axis.delete, bars.y_axis.majorGridlines = True, None
    bars.series[0].graphicalProperties.solidFill = "0E9F6E"
    bars.series[0].dLbls = DataLabelList(showVal=True, showCatName=False, showSerName=False, showLegendKey=False, showPercent=False)
    bars.width, bars.height = 15, 7
    rep.add_chart(bars, "F5")

    months = AreaChart()
    months.title, months.legend = "Calls for help by month", None
    months.add_data(Reference(data, min_col=3, min_row=6, max_row=6), titles_from_data=False)
    months.set_categories(Reference(data, min_col=1, min_row=6, max_row=6))
    months.series[0].graphicalProperties.solidFill = "B7E4CF"
    months.series[0].graphicalProperties.line.solidFill = "0E9F6E"
    months.width, months.height = 15, 5.5
    rep.add_chart(months, "F20")

    rate = LineChart()
    rate.title, rate.legend = "Answered while staffed, by month", None
    rate.add_data(Reference(data, min_col=6, min_row=6, max_row=6), titles_from_data=False)
    rate.set_categories(Reference(data, min_col=1, min_row=6, max_row=6))
    rate.y_axis.numFmt = "0%"
    rate.series[0].graphicalProperties.line.solidFill = "0E9F6E"
    rate.series[0].graphicalProperties.line.width = 28575
    rate.series[0].smooth = False
    rate.width, rate.height = 15, 5.5
    rep.add_chart(rate, "F32")

    for ws in wb.worksheets:
        printable(ws, one_page=ws is rep)
    wb.properties.title = "{{question}}"
    wb.properties.creator = "The Fourth Sheet"
    wb.save(OUT / "calls.xlsx")


# ------------------------------------------------------------------ PowerPoint

class Template(pptkit.Deck):
    def footers(self):
        n = len(self.p.slides)
        for i, s in enumerate(self.p.slides, 1):
            y = pptkit.H - Inches(0.4)
            self.text(s, pptkit.L, y, Inches(3.2), Inches(0.3), "{{business}}", 9, colour=MUTED)
            self.text(s, Inches(3.4), y, Inches(6.5), Inches(0.3), "{{generated}} · Sample data", 9, colour=MUTED, align=2)
            self.text(s, Inches(9.3), y, Inches(3.5), Inches(0.3), f"Slide {i} of {n}", 9, colour=MUTED, align=3)


def tag(slide, names):
    """Name each chart on the slide (its alt text, "chart:<name>") and let its axes scale to whatever data goes in."""
    frames = [sh for sh in slide.shapes if getattr(sh, "has_chart", False) and sh.has_chart]
    for sh, name in zip(frames, names):
        sh._element.xpath("./p:nvGraphicFramePr/p:cNvPr")[0].set("descr", f"chart:{name}")
        cs = sh.chart._chartSpace
        for sc in cs.iter(qn("c:scaling")):
            for t in ("c:min", "c:max"):
                for e in sc.findall(qn(t)):
                    sc.remove(e)
        for ax in cs.iter(qn("c:valAx")):
            for e in ax.findall(qn("c:majorUnit")):
                ax.remove(e)
        for ser in cs.iter(qn("c:ser")):            # one colour for every bar, whatever the count
            for e in ser.findall(qn("c:dPt")):
                ser.remove(e)
        if name == "hours":                         # first time of day at the top
            cat = next(cs.iter(qn("c:catAx")))
            cat.find(qn("c:scaling")).find(qn("c:orientation")).set("val", "maxMin")
            ser = next(cs.iter(qn("c:ser")))
            sp = ser.find(qn("c:spPr"))
            if sp is None:                          # the schema's order: idx, order, tx, spPr, ... (PowerPoint insists)
                sp = ser.makeelement(qn("c:spPr"), {})
                (ser.find(qn("c:tx")) if ser.find(qn("c:tx")) is not None else ser.find(qn("c:order"))).addnext(sp)
            sp.clear()
            fill = sp.makeelement(qn("a:solidFill"), {})
            fill.append(fill.makeelement(qn("a:srgbClr"), {"val": pptkit.GREEN}))
            sp.append(fill)


# Boxes are sized by the length of their text, so each is laid out with stand-in text as long as the longest it will hold,
# then the stand-in is swapped for its slot.
STAND_IN = {"§intro§": ("Period: {{period_label}}. In short: {{answer}} {{intro}}", 330),
            "§shows§": ("{{shows}}", 250), "§action§": ("{{action}}", 250)}


def stand_in(key):
    return (key + " lorem") * (STAND_IN[key][1] // (len(key) + 6))


def swap_stand_ins(deck):
    for s in deck.p.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    for r in para.runs:
                        for key, (slot, _) in STAND_IN.items():
                            if r.text.startswith(key):
                                r.text = slot


def pptx():
    d = Template("{{business}}", "{{question}}", "calls.pptx", "")
    d.title_slide(stand_in("§intro§"), ["{{status_lines}}"])
    labels = [f"{{{{x{i}}}}}" for i in range(7)]                 # placeholder data: the fill replaces it
    s = d.combo_slide("{{question}}", kpis=[(f"{{{{kpi{i}.label}}}}", f"{{{{kpi{i}.value}}}}", f"{{{{kpi{i}.sub}}}}") for i in range(1, 5)],
                      chart={"kind": "bars", "labels": labels, "values": [1] * 7, "fmt": "int", "series_name": "Calls", "keep_order": True},
                      blocks=[("What it shows", stand_in("§shows§")), ("What you'd do about it", stand_in("§action§"))],
                      sub="When the calls come in, {{period_label}}")
    tag(s, ["hours"])
    s, top = d.slide("Calls for help by month", "{{months_note}}")
    h = (pptkit.H - top - Inches(1.0)) / 2
    d.text(s, pptkit.L, top, Inches(6), Inches(0.3), "CALLS", 10, bold=True, colour=MUTED)
    d._area_chart(s, pptkit.L, top + Inches(0.3), pptkit.CONTENT_W, h - Inches(0.35), labels, [1] * 7, "int", "Calls")
    d.text(s, pptkit.L, top + h, Inches(6), Inches(0.3), "ANSWERED WHILE STAFFED", 10, bold=True, colour=MUTED)
    d._area_chart(s, pptkit.L, top + h + Inches(0.3), pptkit.CONTENT_W, h - Inches(0.35), labels, [1] * 7, "pct1", "Answered while staffed %")
    tag(s, ["months_calls", "months_rate"])
    swap_stand_ins(d)
    d.save(OUT / "calls.pptx")


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    xlsx()
    pptx()
    print(f"wrote {OUT / 'calls.xlsx'} and {OUT / 'calls.pptx'}")
