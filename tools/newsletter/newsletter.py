#!/usr/bin/env python3
"""
newsletter.py — the weekly newsletter for The Fourth Sheet: a review spreadsheet in, a blog post out.

WHO IT'S WRITTEN FOR
  Your customers: owners and managers of small businesses and not-for-profits who aren't accountants or data people.
  Every story has to answer "what does this mean for me, and by when?". The post is built to show that:
    - "Dates to hold" box at the top: every story with a clock on it, in one place.
    - A LEAD story (one per issue) with a picture: the thing you most want them to understand this week.
    - EVENTS shown as Who / What / When / Where / So what.
    - NEWS grouped by who it lands on: businesses, not-for-profits, or both.

HOW IT WORKS
  1. Each week has ONE Excel file:  tools/newsletter/issues/<YYYY>-W<NN>/review.xlsx
       - "Items" sheet: one row per story. Claude fills the research columns; YOU tick (Include = Yes / No),
         add or delete rows, set the Type (Lead / News / Event) and write the "My take" column in your own words.
       - "Settings" sheet: newsletter name, issue title, week, status (Draft / Published), your opening note and sign-off.
  2. Only rows marked Include = Yes make it into the post. Published text is yours or a plain factual label:
       Headline, Source, Link, Date, Where, Act by, image alt text and credit are labels you can edit;
       "My take", the opening note and the sign-off are your voice. The "Facts" and "Owner's worry" columns are
       research notes for you and are NEVER published.
  3. Pictures: put the image file in newsletter/images/ and type its file name in the Image column (or paste an https
       address). Fill in Image alt (what it shows) and Image credit (who made it / licence). Only use pictures you made
       or have the right to use. The build checks the file exists.

COMMANDS (run from the repo root; Windows: py instead of python3)
    python3 tools/newsletter/newsletter.py new 2026-W44        # blank review file for a week
    python3 tools/newsletter/newsletter.py preview 2026-W43    # build it and open it in your browser. Nothing is published.
    python3 tools/newsletter/newsletter.py build 2026-W43      # render that week's post + refresh the newsletter index
    python3 tools/newsletter/newsletter.py build-all           # rebuild every week (after a style change)

STATUS
  Draft      -> post is written to newsletter/<week>.html with "noindex", a DRAFT banner, and is NOT listed on the index.
  Published  -> the build STOPS unless the opening note, sign-off and a "My take" for every included item are filled in,
                links are valid, and every picture exists and has alt text. The post loses the banner and noindex and is
                listed on newsletter/index.html.
  Nothing goes live until the file is committed and pushed to main (the site publishes on every push).
"""

import html
import re
import sys
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

REPO = Path(__file__).resolve().parents[2]
ISSUES = REPO / "tools" / "newsletter" / "issues"      # the review files (kept in git, not published: tools/ is excluded)
OUT = REPO / "newsletter"                               # the published posts
IMAGES = OUT / "images"                                 # pictures used in posts
DOMAIN = "https://www.thefourthsheet.com.au"

# ---- the columns of the Items sheet, in order. The header text is what the build looks for. ----
ITEM_COLS = [
    ("Include", 9),                          # Yes / No dropdown: the tick and flick
    ("Type", 9),                             # Lead (one per issue) / News / Event
    ("Audience", 11),                        # SME / NFP / Both: who it lands on (News section)
    ("Lens", 13),                            # Finance / Operations / Both / Data craft: shown as a small tag
    ("Headline", 44),                        # published: plain label for the story (WHAT)
    ("Source", 24),                          # published: who said it / who runs the event (WHO)
    ("Link", 40),                            # published
    ("Date", 14),                            # published: when the source came out, or the event date (WHEN)
    ("Where", 20),                           # published (events): venue or online (WHERE)
    ("Act by (published)", 22),              # published: the date or trigger that makes it urgent
    ("Image", 26),                           # file name in newsletter/images/, or an https address
    ("Image alt", 30),                       # published: what the picture shows (accessibility)
    ("Image credit", 22),                    # published: who made it / licence
    ("Facts (not published)", 60),           # research notes: what is actually true, with the gaps flagged
    ("Owner's worry (not published)", 50),   # the question an owner is really asking; a prompt for your take
    ("My take (published)", 70),             # YOUR words: the SO WHAT. Blank + Include=Yes means the item is not ready.
    ("Added by", 10),                        # Claude / Nathan
]
COL = {name: i for i, (name, _) in enumerate(ITEM_COLS)}
TYPES = ["Lead", "News", "Event"]
AUDIENCES = ["SME", "NFP", "Both"]
LENSES = ["Finance", "Operations", "Both", "Data craft"]
AUDIENCE_TITLES = {"SME": "If you run a business", "NFP": "If you run a not-for-profit", "Both": "Whichever you run"}

# ---- the Settings sheet: key, default, note ----
SETTINGS = [
    ("Newsletter name", "The Fourth Read", "Working title. Change it here and every post and the index change with it."),
    ("Tagline", "Data, operations and the numbers underneath. Weekly.", "Shown under the name. Yours to change."),
    ("Issue title", "", "Optional. A line for this week, in your words. Blank = just the newsletter name."),
    ("Week label", "", "e.g. Week of 19 October 2026"),
    ("Status", "Draft", "Draft or Published. Published is blocked until your words are in."),
    ("Opening note", "", "YOURS. A few lines on what you noticed this week and why your owners should care."),
    ("Sign-off", "", "YOURS. How you close it."),
]


# ------------------------------------------------------------------------------------------ the review file
def week_dates(week):
    """'2026-W43' -> (Monday, Sunday) of that ISO week."""
    m = re.fullmatch(r"(\d{4})-W(\d{2})", week)
    if not m:
        sys.exit(f"Week must look like 2026-W43, not {week!r}")
    monday = date.fromisocalendar(int(m.group(1)), int(m.group(2)), 1)
    return monday, monday + timedelta(days=6)


def new_review(week, items=None, settings=None):
    """Create issues/<week>/review.xlsx. `items` is a list of dicts keyed by the ITEM_COLS names (Claude's research);
    `settings` overrides the Settings defaults."""
    monday, _ = week_dates(week)
    folder = ISSUES / week
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / "review.xlsx"
    if path.exists():
        sys.exit(f"{path} already exists. Not overwriting your ticks.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Items"
    names = [c for c, _ in ITEM_COLS]
    ws.append(names)
    for col, (_, width) in enumerate(ITEM_COLS, start=1):
        c = ws.cell(1, col)
        ws.column_dimensions[c.column_letter].width = width
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F7A5D")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "F2"

    rows = list(items or [])
    for r in rows:
        ws.append([r.get(n, "") for n in names])
    SPARE = 12                                    # empty rows with the dropdowns already on, ready for new stories
    for _ in range(SPARE):
        r = {"Include": "No", "Type": "News", "Audience": "Both", "Added by": "Nathan"}
        ws.append([r.get(n, "") for n in names])
    last = ws.max_row
    for row in ws.iter_rows(min_row=2, max_row=last):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)

    reach = last + 200                            # dropdowns run well past the last row, so new rows just work
    for col_name, options in [("Include", ["Yes", "No"]), ("Type", TYPES), ("Audience", AUDIENCES), ("Lens", LENSES)]:
        letter = ws.cell(1, COL[col_name] + 1).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=(col_name == "Lens"))
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{reach}")
    end = ws.cell(1, len(ITEM_COLS)).column_letter
    head = ws.cell(1, COL["Headline"] + 1).column_letter
    ws.conditional_formatting.add(f"A2:{end}{reach}", FormulaRule(formula=['$A2="Yes"'], fill=PatternFill("solid", bgColor="E3F4E8")))
    ws.conditional_formatting.add(f"A2:{end}{reach}", FormulaRule(formula=[f'AND($A2="No",${head}2<>"")'], fill=PatternFill("solid", bgColor="EEF2EE")))

    st = wb.create_sheet("Settings")
    st.append(["Setting", "Value", "Note"])
    for c in st[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F7A5D")
    st.column_dimensions["A"].width = 18
    st.column_dimensions["B"].width = 80
    st.column_dimensions["C"].width = 70
    over = settings or {}
    status_row = None
    for key, default, note in SETTINGS:
        default = over.get(key, default)
        if key == "Week label" and not default:
            default = f"Week of {monday.day} {monday:%B %Y}"
        st.append([key, default, note])
        if key == "Status":
            status_row = st.max_row
    for row in st.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    dv = DataValidation(type="list", formula1='"Draft,Published"', allow_blank=False)
    st.add_data_validation(dv)
    dv.add(f"B{status_row}")

    wb.save(path)
    print(f"created {path.relative_to(REPO)}  ({len(rows)} researched items + {SPARE} spare rows)")
    return path


def read_review(week):
    path = ISSUES / week / "review.xlsx"
    if not path.exists():
        sys.exit(f"No review file at {path}. Make one with:  python3 tools/newsletter/newsletter.py new {week}")
    wb = load_workbook(path, data_only=True)
    settings = {}
    for key, value, _ in wb["Settings"].iter_rows(min_row=2, values_only=True):
        if key:
            settings[str(key).strip()] = ("" if value is None else str(value).strip())
    header = [c.value for c in wb["Items"][1]]
    items = []
    for row in wb["Items"].iter_rows(min_row=2, values_only=True):
        rec = {h: ("" if v is None else (v.strftime("%-d %B %Y") if hasattr(v, "strftime") else str(v).strip()))
               for h, v in zip(header, row)}
        if rec.get("Headline"):                   # a row with no headline is an empty spare row
            items.append(rec)
    return settings, items


# ------------------------------------------------------------------------------------------ the post
def esc(s):
    return html.escape(s, quote=True)


def paragraphs(text):
    """Blank line = new paragraph. Nathan's words go in as typed (escaped), nothing added."""
    return "\n".join(f"<p>{esc(p.strip()).replace(chr(10), '<br>')}</p>" for p in re.split(r"\n\s*\n", text) if p.strip())


def safe_link(url):
    return url if re.match(r"https?://", url) else ""


def image_src(ref):
    """'file.svg' -> '/newsletter/images/file.svg' if the file exists; an https address passes through; else ''."""
    if not ref:
        return ""
    if re.match(r"https://", ref):
        return ref
    return f"/newsletter/images/{ref}" if (IMAGES / ref).is_file() else ""


NAV = """<header class="site-header">
      <div class="container nav-bar">
        <a class="brand" href="/"><img class="brand-mark" src="/assets/brand/mark-4th.svg?v=sheet" alt="" width="32" height="32" /><span>The Fourth Sheet</span></a>
        <button class="nav-toggle" aria-label="Open menu" aria-expanded="false" aria-controls="primary-nav">☰</button>
        <nav class="nav-links" id="primary-nav" aria-label="Primary navigation">
          <a href="/sme.html">Businesses</a>
          <a href="/not-for-profit.html">Not-for-profits</a>
          <a href="/examples.html">Examples</a>
          <a href="/services.html">Services</a>
          <a href="/about.html">About</a>
          <a href="/contact.html">Contact</a>
          <a class="button primary nav-book" href="/contact.html?topic=numbers-check" data-track="numbers-check">Ask for a free 20-minute numbers check</a>
        </nav>
      </div>
    </header>"""

FOOTER = """<footer class="site-footer">
      <div class="container footer-content">
        <nav class="footer-links" aria-label="Site">
          <a href="/sme.html">Businesses</a><a href="/not-for-profit.html">Not-for-profits</a><a href="/examples.html">Examples</a><a href="/services.html">Services</a><a href="/about.html">About</a><a href="/numbers-explained.html">Numbers explained</a><a href="/contact.html">Contact</a>
        </nav>
        <p>© <span id="year">2026</span> The Fourth Sheet · Nathan Booth · Brisbane, Australia · ABN 31 839 620 153</p>
      </div>
    </footer>"""

# small extra styles for the newsletter only (the site's styles.css is untouched), using the site's own colour variables
NL_CSS = """<style>
      main a:not(.button) { color: var(--accent); }
      .nl-clock { background: var(--accent-soft); border: 1px solid var(--accent-tint-border); border-radius: var(--r-md); padding: var(--sp-5); margin-top: var(--sp-6); }
      .nl-clock h2 { margin-top: 0; font-size: var(--fs-lg); }
      .nl-clock ul { margin: 0; padding-left: var(--sp-5); }
      .nl-clock li { margin: var(--sp-2) 0; }
      .nl-lead { display: grid; gap: var(--sp-5); }
      .nl-lead img, .nl-card img { width: 100%; height: auto; border-radius: var(--r-md); border: 1px solid var(--border); background: var(--paper); }
      .nl-credit { font-size: var(--fs-xs); color: var(--muted); margin: var(--sp-1) 0 0; }
      .nl-chip { display: inline-block; font-size: var(--fs-xs); border-radius: var(--r-pill); padding: 2px 10px; margin-right: var(--sp-2); background: var(--accent-soft); color: var(--accent); }
      .nl-chip.urgent { background: var(--bad-soft); color: var(--bad); font-weight: 700; }
      .nl-facts { display: grid; grid-template-columns: max-content 1fr; gap: var(--sp-1) var(--sp-4); margin: var(--sp-3) 0; }
      .nl-facts dt { font-weight: 700; color: var(--muted); }
      .nl-facts dd { margin: 0; }
      .nl-sowhat { font-weight: 700; margin: var(--sp-4) 0 var(--sp-1); }
    </style>"""


def page(title, description, body, draft=False, canonical=None):
    robots = '<meta name="robots" content="noindex, nofollow" />' if draft else ""
    canon = f'<link rel="canonical" href="{canonical}" />' if canonical and not draft else ""
    return f"""<!DOCTYPE html>
<html lang="en-AU">
  <head>
    <meta charset="UTF-8" />
    <meta name="viewport" content="width=device-width, initial-scale=1.0" />
    <title>{esc(title)} | The Fourth Sheet</title>
    <meta name="description" content="{esc(description)}" />
    {robots}
    {canon}
    <link rel="icon" href="/assets/brand/mark-4th.svg?v=sheet" type="image/svg+xml" />
    <meta name="theme-color" content="#2f7a5d" />
    <link rel="preconnect" href="https://fonts.googleapis.com">
    <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
    <link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Roboto:wght@400;500;700&display=swap">
    <link rel="stylesheet" href="/styles.css?v=dev" />
    {NL_CSS}
  </head>
  <body>
    {NAV}
    <main id="main">
{body}
    </main>
    {FOOTER}
    <script src="/script.js?v=dev"></script>
  </body>
</html>
"""


def problems_for_publish(settings, included):
    """Everything that must be in order before a post can go out as Published."""
    p = []
    if not settings.get("Opening note"):
        p.append("Settings: Opening note is empty")
    if not settings.get("Sign-off"):
        p.append("Settings: Sign-off is empty")
    if sum(1 for i in included if i.get("Type") == "Lead") > 1:
        p.append("More than one Lead ticked: only one per issue")
    for it in included:
        h = it["Headline"]
        if not it.get("My take (published)"):
            p.append(f"No 'My take' yet: {h}")
        if not safe_link(it.get("Link", "")):
            p.append(f"Missing or invalid link: {h}")
        if it.get("Type") == "Event" and not (it.get("Date") and it.get("Where")):
            p.append(f"Event needs a Date and a Where: {h}")
        if it.get("Image"):
            if not image_src(it["Image"]):
                p.append(f"Picture not found in newsletter/images/: {it['Image']} ({h})")
            if not it.get("Image alt"):
                p.append(f"Picture has no alt text: {h}")
    if not included:
        p.append("No items ticked Yes")
    return p


def picture(it):
    src = image_src(it.get("Image", ""))
    if not src:
        return ""
    credit = f'<p class="nl-credit">{esc(it["Image credit"])}</p>' if it.get("Image credit") else ""
    return f'<figure style="margin:0"><img src="{esc(src)}" alt="{esc(it.get("Image alt", ""))}" loading="lazy" />{credit}</figure>'


def chips(it):
    out = []
    if it.get("Lens"):
        out.append(f'<span class="nl-chip">{esc(it["Lens"])}</span>')
    if it.get("Act by (published)"):
        out.append(f'<span class="nl-chip urgent">Act by: {esc(it["Act by (published)"])}</span>')
    return "".join(out)


def title_link(it):
    link = safe_link(it.get("Link", ""))
    head = esc(it["Headline"])
    return f'<a href="{esc(link)}" rel="noopener" target="_blank">{head}</a>' if link else head


def build(week):
    settings, items = read_review(week)
    monday, _ = week_dates(week)
    name = settings.get("Newsletter name") or "The Fourth Read"
    status = settings.get("Status", "Draft")
    draft = status != "Published"
    included = [i for i in items if i.get("Include", "").lower() == "yes"]
    todo = problems_for_publish(settings, included)
    if not draft and todo:
        print("NOT PUBLISHED. Status says Published but these still need sorting:")
        for t in todo:
            print("  -", t)
        sys.exit(1)

    week_label = settings.get("Week label") or f"Week of {monday.day} {monday:%B %Y}"
    issue_title = settings.get("Issue title", "")
    placeholder = '<span class="tbc">[Nathan to write]</span>'
    take_html = lambda it: paragraphs(it["My take (published)"]) if it.get("My take (published)") else f"<p>{placeholder}</p>"

    parts = []
    banner = ""
    if draft:
        banner = (f'<p class="note-card"><strong>DRAFT, not published.</strong> {len(included)} stories ticked. '
                  f'Still to sort: {len(todo)} item(s). This page is hidden from search engines and not linked anywhere.</p>')
    opening = settings.get("Opening note", "")

    # --- "Dates to hold": anything with an Act by, plus every event, in the order they're listed in the file
    clock = [i for i in included if i.get("Act by (published)") or i.get("Type") == "Event"]
    clock_html = ""
    if clock:
        lis = "\n".join(
            f'        <li><strong>{esc(i.get("Act by (published)") or i.get("Date") or "")}</strong> · {esc(i["Headline"])}</li>' for i in clock)
        clock_html = f'<div class="nl-clock"><h2>Dates to hold</h2>\n      <ul>\n{lis}\n      </ul></div>'

    parts.append(f"""      <section class="section page-top section-alt">
        <div class="container">
          <div class="section-heading">
            <p class="eyebrow">{esc(name)} · {esc(week_label)}</p>
            <h1>{esc(issue_title) if issue_title else esc(name)}</h1>
            <p class="lead">{esc(settings.get('Tagline', ''))}</p>
          </div>
          {banner}
          <div>{paragraphs(opening) if opening else '<p>' + placeholder + '</p>'}</div>
          {clock_html}
        </div>
      </section>""")

    # --- the lead
    for it in [i for i in included if i.get("Type") == "Lead"][:1]:
        parts.append(f"""      <section class="section">
        <div class="container nl-lead">
          <p class="eyebrow">This week's lead</p>
          <h2>{title_link(it)}</h2>
          <p>{chips(it)}</p>
          {picture(it)}
          {take_html(it)}
        </div>
      </section>""")

    # --- events: who / what / when / where / so what
    events = [i for i in included if i.get("Type") == "Event"]
    if events:
        cards = []
        for it in events:
            cards.append(f"""            <article class="info-card nl-card">
              <h3>{title_link(it)}</h3>
              <p>{chips(it)}</p>
              {picture(it)}
              <dl class="nl-facts">
                <dt>What</dt><dd>{esc(it["Headline"])}</dd>
                <dt>Who</dt><dd>{esc(it.get("Source", ""))}</dd>
                <dt>When</dt><dd>{esc(it.get("Date", ""))}</dd>
                <dt>Where</dt><dd>{esc(it.get("Where", ""))}</dd>
              </dl>
              <p class="nl-sowhat">So what</p>
              {take_html(it)}
            </article>""")
        parts.append(f"""      <section class="section section-alt">
        <div class="container">
          <div class="section-heading"><h2>Events worth your time</h2></div>
          <div class="card-grid card-grid--two">
{chr(10).join(cards)}
          </div>
        </div>
      </section>""")

    # --- news, grouped by who it lands on
    for aud in AUDIENCES:
        group = [i for i in included if i.get("Type", "News") in ("News", "") and (i.get("Audience") or "Both") == aud]
        if not group:
            continue
        cards = []
        for it in group:
            meta = " · ".join(x for x in [esc(it.get("Source", "")), esc(it.get("Date", ""))] if x)
            cards.append(f"""            <article class="info-card nl-card">
              <h3>{title_link(it)}</h3>
              <p>{chips(it)}</p>
              <p class="muted">{meta}</p>
              {picture(it)}
              <p class="nl-sowhat">So what</p>
              {take_html(it)}
            </article>""")
        parts.append(f"""      <section class="section">
        <div class="container">
          <div class="section-heading"><h2>{esc(AUDIENCE_TITLES[aud])}</h2></div>
          <div class="card-grid card-grid--two">
{chr(10).join(cards)}
          </div>
        </div>
      </section>""")

    signoff = settings.get("Sign-off", "")
    parts.append(f"""      <section class="section section-alt">
        <div class="container">
          {paragraphs(signoff) if signoff else '<p>' + placeholder + '</p>'}
          <p><a class="button primary" href="/contact.html?topic=numbers-check" data-track="numbers-check">Ask for a free 20-minute numbers check</a></p>
          <p><a href="/newsletter/">All issues of {esc(name)}</a></p>
        </div>
      </section>""")

    OUT.mkdir(exist_ok=True)
    slug = week.lower()
    desc = f"{name}, {week_label}: what's changing for small businesses and not-for-profits, and what to do about it."
    (OUT / f"{slug}.html").write_text(
        page(f"{name}: {week_label}", desc, "\n".join(parts), draft, f"{DOMAIN}/newsletter/{slug}.html"), encoding="utf-8")
    print(f"built newsletter/{slug}.html  [{status}]  {len(included)} stories ticked, {len(todo)} still to sort")
    build_index()


def build_index():
    """newsletter/index.html: lists Published issues only, newest first."""
    published = []
    name, tagline = "The Fourth Read", ""
    for review in sorted(ISSUES.glob("*/review.xlsx"), reverse=True):
        week = review.parent.name
        settings, items = read_review(week)
        name = settings.get("Newsletter name") or name
        tagline = settings.get("Tagline") or tagline
        if settings.get("Status") == "Published":
            n = sum(1 for i in items if i.get("Include", "").lower() == "yes")
            published.append((week, settings.get("Week label", week), settings.get("Issue title", ""), n))
    if published:
        rows = "\n".join(
            f'          <li><a href="/newsletter/{w.lower()}.html"><strong>{esc(label)}</strong></a>'
            f'{" · " + esc(t) if t else ""} <span class="muted">({n} stories)</span></li>'
            for w, label, t, n in published)
        listing = f"<ul>\n{rows}\n        </ul>"
    else:
        listing = "<p>The first issue isn't out yet.</p>"
    body = f"""      <section class="section page-top section-alt">
        <div class="container">
          <div class="section-heading">
            <p class="eyebrow">Newsletter</p>
            <h1>{esc(name)}</h1>
            <p class="lead">{esc(tagline)}</p>
          </div>
          {listing}
        </div>
      </section>"""
    OUT.mkdir(exist_ok=True)
    (OUT / "index.html").write_text(
        page(name, f"{name}: weekly news for small businesses and not-for-profits, read through the numbers.", body,
             draft=not published, canonical=f"{DOMAIN}/newsletter/"), encoding="utf-8")


# ------------------------------------------------------------------------------------------ preview
def preview(week):
    """Build the week, serve the site folder on your own machine and open the page. Drafts show here whatever the
    Status; nothing is committed or published. Stop with Ctrl-C."""
    import functools, http.server, threading, webbrowser
    build(week)
    handler = functools.partial(http.server.SimpleHTTPRequestHandler, directory=str(REPO))
    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), handler)
    url = f"http://127.0.0.1:{server.server_address[1]}/newsletter/{week.lower()}.html"
    print(f"previewing {url}   (Ctrl-C to stop)")
    threading.Timer(0.5, lambda: webbrowser.open(url)).start()
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("stopped")


# ------------------------------------------------------------------------------------------ command line
if __name__ == "__main__":
    args = sys.argv[1:]
    if len(args) == 2 and args[0] == "new":
        new_review(args[1])
    elif len(args) == 2 and args[0] == "build":
        build(args[1])
    elif len(args) == 2 and args[0] == "preview":
        preview(args[1])
    elif args == ["build-all"]:
        for r in sorted(ISSUES.glob("*/review.xlsx")):
            build(r.parent.name)
    else:
        print(__doc__)
        sys.exit(2)
